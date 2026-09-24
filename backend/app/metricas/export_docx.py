"""Reporte de métricas en Word — el reporte trimestral que hoy se arma a
mano: mismas cifras que el dashboard, para el período elegido."""

import io
from datetime import date

from docx import Document

from app.metricas.schemas import MetricasOut

TIPO_LABELS = {
    "candidato": "Candidato",
    "ingresado": "Ingresado",
    "herramienta_tercero": "Herramienta de tercero",
}
SEMAFORO_LABELS = {
    "solido": "Sólido",
    "moderado": "Moderado",
    "vulnerable": "Vulnerable",
    "critico": "Crítico",
}
SEVERIDAD_LABELS = {"baja": "Baja", "media": "Media", "alta": "Alta", "critica": "Crítica"}


def _pct(value: float) -> str:
    return f"{value * 100:.1f}%"


def _table(document: Document, header: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    for i, title in enumerate(header):
        table.rows[0].cells[i].text = title
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value


def build_metricas_report(
    metricas: MetricasOut, desde: date | None, hasta: date | None, tipo: str | None
) -> bytes:
    document = Document()
    document.add_heading("AI-SEC Control — Reporte de métricas", level=0)

    periodo = f"{desde or 'inicio'} a {hasta or 'hoy'}"
    document.add_paragraph(f"Período: {periodo}")
    document.add_paragraph(f"Tipo de caso: {TIPO_LABELS.get(tipo, 'Todos') if tipo else 'Todos'}")

    pi, an, mo, pe = (
        metricas.proyectos_ingresados,
        metricas.analisis,
        metricas.modelados,
        metricas.pentesting,
    )

    document.add_heading("Resumen", level=1)
    _table(
        document,
        ["Métrica", "Cantidad"],
        [
            ["Proyectos ingresados", str(pi.total)],
            ["Casos analizados con IA", str(an.casos_analizados)],
            ["Modelados de amenazas realizados", str(mo.cantidad)],
            ["Pentestings realizados", str(pe.total)],
        ],
    )

    document.add_heading("Proyectos ingresados", level=1)
    _table(
        document,
        ["Tipo", "Ingresados", "Analizados con IA"],
        [[TIPO_LABELS[t], str(pi.por_tipo[t]), str(an.por_tipo[t])] for t in TIPO_LABELS],
    )
    if pi.por_mes:
        document.add_paragraph("")
        _table(document, ["Mes", "Proyectos"], [[m.periodo, str(m.cantidad)] for m in pi.por_mes])

    document.add_heading("Modelados de amenazas y riesgo resultante", level=1)
    if mo.cantidad:
        document.add_paragraph(
            f"Cumplimiento promedio: {_pct(mo.promedio_compliance_pct or 0)} · "
            f"Riesgo residual promedio: {_pct(mo.promedio_residual_pct or 0)}"
        )
        _table(
            document,
            ["Proyecto", "Empresa", "Tipo", "Cumplimiento", "Riesgo residual", "Brechas críticas", "Semáforo"],
            [
                [
                    c.nombre_proyecto,
                    c.empresa_responsable,
                    TIPO_LABELS[c.tipo],
                    _pct(c.compliance_pct),
                    _pct(c.residual_pct),
                    str(c.brechas_criticas),
                    SEMAFORO_LABELS[c.semaforo],
                ]
                for c in mo.casos
            ],
        )
    else:
        document.add_paragraph("Ningún caso del período tiene preguntas respondidas.")

    document.add_heading("Pentesting", level=1)
    if pe.total:
        document.add_paragraph(f"{pe.total} resultado(s) en {pe.casos_con_pentest} caso(s).")
        _table(document, ["Herramienta", "Resultados"], [[h.nombre, str(h.cantidad)] for h in pe.por_herramienta])
        document.add_paragraph("")
        _table(
            document,
            ["Severidad", "Resultados"],
            [[SEVERIDAD_LABELS[s], str(pe.por_severidad[s])] for s in SEVERIDAD_LABELS],
        )
    else:
        document.add_paragraph("Sin pentestings en el período.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
