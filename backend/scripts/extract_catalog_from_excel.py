"""Extrae el catálogo del Excel PLOT4AI a un JSON — el insumo de una migración
de datos (alembic/versions/8c55274c7ed7_* carga catalogo_maestro_v1.json).
La app nunca lee el .xlsx en runtime.

El JSON de cada migración es un snapshot congelado: este script se niega a
pisar un archivo existente. Para cambiar el catálogo, generar una versión
nueva y escribir una migración que la aplique:

    uv run python scripts/extract_catalog_from_excel.py --output alembic/data/catalogo_maestro_v2.json
"""

import argparse
import json
import sys
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent.parent
EXCEL_PATH = REPO_ROOT / "AISEC_PLOT4AI_Final -05-2026-Proveedor.xlsx"
BACKEND_DIR = Path(__file__).resolve().parent.parent
DEFAULT_OUTPUT = BACKEND_DIR / "alembic" / "data" / "catalogo_maestro_v1.json"

CATALOGO_SHEET = "📋 Catálogo Maestro"
PARAMETROS_SHEET = "⚙️ Parámetros"

TIER_MAP = {"🔴 ×3": "critico", "🟡 ×2": "alto", "⚪ ×1": "estandar"}
TIER_MULTIPLICADOR = {"critico": 3, "alto": 2, "estandar": 1}
TIPO_MAP = {"Control ✅": "control", "Riesgo ⚠️": "riesgo"}


def _parse_tipo(raw: str) -> str:
    try:
        return TIPO_MAP[raw]
    except KeyError as exc:
        raise ValueError(f"Tipo desconocido: {raw!r}") from exc


def _parse_tier(raw: str) -> str:
    try:
        return TIER_MAP[raw]
    except KeyError as exc:
        raise ValueError(f"Tier desconocido: {raw!r}") from exc


def _parse_polaridad(raw: str) -> str:
    if raw.startswith("+1"):
        return "positiva"
    if raw.startswith("-1"):
        return "negativa"
    raise ValueError(f"Polaridad desconocida: {raw!r}")


def extract_dominios(wb: openpyxl.Workbook) -> list[dict]:
    ws = wb[PARAMETROS_SHEET]
    dominios = []
    for row in range(7, 15):  # rows 7..14 per the "PESOS POR DOMINIO" table
        nombre = ws.cell(row=row, column=2).value  # B
        peso = ws.cell(row=row, column=3).value  # C
        total_qs = ws.cell(row=row, column=4).value  # D
        if nombre is None:
            continue
        dominios.append(
            {
                "nombre": nombre,
                "peso": float(peso),
                "total_preguntas": int(total_qs),
            }
        )
    return dominios


def _domain_row_lookup(wb: openpyxl.Workbook, dominio_nombre: str) -> dict[int, dict]:
    """Maps question '#' -> {polaridad, explicacion_control_es} for one domain sheet."""
    ws = wb[dominio_nombre]
    lookup: dict[int, dict] = {}
    for row in range(4, ws.max_row + 1):
        numero = ws.cell(row=row, column=1).value  # A
        if not isinstance(numero, (int, float)):
            continue  # below the question rows sit summary/footer rows, not data
        polaridad_raw = ws.cell(row=row, column=5).value  # E
        explicacion = ws.cell(row=row, column=21).value  # U
        lookup[int(numero)] = {
            "polaridad": _parse_polaridad(polaridad_raw),
            "explicacion_control_es": explicacion,
        }
    return lookup


def extract_preguntas(wb: openpyxl.Workbook, codigo_by_nombre: dict[str, str]) -> list[dict]:
    ws = wb[CATALOGO_SHEET]
    domain_lookups: dict[str, dict[int, dict]] = {}
    preguntas = []

    for row in range(4, ws.max_row + 1):
        id_ = ws.cell(row=row, column=1).value  # A
        if id_ is None:
            continue
        codigo = ws.cell(row=row, column=2).value  # B
        dominio_nombre = ws.cell(row=row, column=3).value  # C
        numero = int(ws.cell(row=row, column=4).value)  # D "#"
        texto_es = ws.cell(row=row, column=5).value  # E
        texto_en = ws.cell(row=row, column=6).value  # F
        tipo_raw = ws.cell(row=row, column=7).value  # G
        tier_raw = ws.cell(row=row, column=8).value  # H
        impacto_primario = ws.cell(row=row, column=9).value  # I
        referencias_regulatorias = ws.cell(row=row, column=10).value  # J
        justificacion_tier = ws.cell(row=row, column=11).value  # K

        if dominio_nombre not in domain_lookups:
            domain_lookups[dominio_nombre] = _domain_row_lookup(wb, dominio_nombre)
        extra = domain_lookups[dominio_nombre][numero]

        tier = _parse_tier(tier_raw)
        preguntas.append(
            {
                "id_legado": id_,
                "dominio_codigo": codigo,
                "numero": numero,
                "texto_es": texto_es,
                "texto_en": texto_en,
                "tipo": _parse_tipo(tipo_raw),
                "polaridad": extra["polaridad"],
                "tier": tier,
                "multiplicador": TIER_MULTIPLICADOR[tier],
                "impacto_primario": impacto_primario,
                "referencias_regulatorias": referencias_regulatorias,
                "justificacion_tier": justificacion_tier,
                "explicacion_control_es": extra["explicacion_control_es"],
                # No hay columna en inglés en el Excel fuente para la
                # explicación del control: se deja null como contenido
                # pendiente de traducir (gap de contenido, no de código).
                "explicacion_control_en": None,
            }
        )

    return preguntas


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    output_path = args.output if args.output.is_absolute() else BACKEND_DIR / args.output
    if output_path.exists():
        print(
            f"{output_path} ya existe y es el snapshot de una migración: no se pisa. "
            "Usá --output con un nombre de versión nueva.",
            file=sys.stderr,
        )
        raise SystemExit(1)

    if not EXCEL_PATH.exists():
        print(f"No se encontró el Excel fuente: {EXCEL_PATH}", file=sys.stderr)
        raise SystemExit(1)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    dominios = extract_dominios(wb)
    codigo_by_nombre = {}
    ws_catalogo = wb[CATALOGO_SHEET]
    for row in range(4, ws_catalogo.max_row + 1):
        nombre = ws_catalogo.cell(row=row, column=3).value
        codigo = ws_catalogo.cell(row=row, column=2).value
        if nombre and codigo:
            codigo_by_nombre.setdefault(nombre, codigo)

    for dominio in dominios:
        dominio["codigo"] = codigo_by_nombre[dominio["nombre"]]

    preguntas = extract_preguntas(wb, codigo_by_nombre)

    payload = {"dominios": dominios, "preguntas": preguntas}

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False), encoding="utf-8"
    )
    print(f"Escritas {len(dominios)} dominios y {len(preguntas)} preguntas en {output_path}")


if __name__ == "__main__":
    main()
