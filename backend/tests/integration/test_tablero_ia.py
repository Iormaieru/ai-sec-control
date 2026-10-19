"""Tablero "Seguridad de IA": agregaciones y exportación a Excel y CSV."""

import csv
import datetime as dt
import io

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.alertas.models import (
    Alerta,
    AlertaEvento,
    EquipoTenable,
    EstadoAlerta,
    MetricaIA,
    ModoRegla,
    ReglaAlerta,
    SeveridadAlerta,
)

HOY = dt.date(2026, 10, 19)
RANGO = {"desde": "2026-10-13", "hasta": "2026-10-19"}


def _momento(dias_atras: int, hora: int = 12) -> dt.datetime:
    return dt.datetime.combine(HOY - dt.timedelta(days=dias_atras), dt.time(hora), dt.UTC)


def _datos(db: Session) -> dict[str, ReglaAlerta]:
    """Dos reglas de alerta (alta y crítica) y una de métrica, con eventos en el rango y uno afuera."""
    mcp = db.scalar(select(ReglaAlerta).where(ReglaAlerta.nombre == "Uso de servidores MCP"))
    deepseek = ReglaAlerta(nombre="Subida a DeepSeek", consulta="app like deepseek", severidad=SeveridadAlerta.CRITICA)
    metrica = ReglaAlerta(
        nombre="Subida a IA generativa", consulta="activity eq 'Upload'", severidad=SeveridadAlerta.MEDIA,
        modo=ModoRegla.METRICA,
    )
    db.add_all([deepseek, metrica])
    db.flush()

    def alerta(regla, usuario, app, eventos, equipos, **cambios):
        a = Alerta(
            regla_id=regla.id, app=app, usuario=usuario, severidad=regla.severidad,
            primera_vez=eventos[0][0], ultima_vez=eventos[-1][0], cantidad_eventos=len(eventos),
            equipos=equipos, **cambios,
        )
        db.add(a)
        db.flush()
        for n, (fecha, actividad) in enumerate(eventos):
            db.add(AlertaEvento(alerta_id=a.id, evento_id=f"{a.id}-{n}", fecha=fecha, actividad=actividad,
                                objeto="=HYPERLINK(\"http://x\")" if n == 0 else None, datos={"hostname": equipos[0]}))
        return a

    creada = _momento(5)
    alerta(mcp, "ext@bmexternos.com.ar", "MCP: katalon", [(_momento(5), "CallToolRequest"), (_momento(4), "ListToolsRequest")],
           ["bmn1"], created_at=creada, atendida_at=creada + dt.timedelta(hours=2), estado=EstadoAlerta.EN_REVISION)
    alerta(deepseek, "ana@macro.com.ar", "DeepSeek", [(_momento(1), "Upload"), (_momento(1, 13), "Upload"), (_momento(0), "Post")],
           ["bmn2"], created_at=creada, atendida_at=creada + dt.timedelta(hours=4), cerrada_at=creada + dt.timedelta(hours=10),
           estado=EstadoAlerta.RESUELTA)
    # Fuera del rango: no cuenta.
    alerta(deepseek, "viejo@macro.com.ar", "DeepSeek", [(_momento(40), "Upload")], ["bmn3"])
    db.add_all([
        MetricaIA(regla_id=metrica.id, dia=HOY - dt.timedelta(days=1), app="ChatGPT", usuario="ana@macro.com.ar", actividad="Upload", eventos=10),
        MetricaIA(regla_id=metrica.id, dia=HOY, app="ChatGPT", usuario="beto@macro.com.ar", actividad="Upload", eventos=5),
        MetricaIA(regla_id=metrica.id, dia=HOY - dt.timedelta(days=30), app="Gemini", usuario="beto@macro.com.ar", actividad="Upload", eventos=99),
        EquipoTenable(asset_uuid="t1", hostname="bmn1", con_inventario=True, tiene_agente=True, software_ia=["Python"],
                      actualizado_at=_momento(0)),
    ])
    db.commit()
    return {"mcp": mcp, "deepseek": deepseek, "metrica": metrica}


def test_indicadores_y_series(client: TestClient, auth_headers: dict[str, str], db: Session) -> None:
    _datos(db)

    r = client.get("/metricas/ia", params=RANGO, headers=auth_headers)

    assert r.status_code == 200
    datos = r.json()
    ind = datos["indicadores"]
    assert (ind["alertas_abiertas"], ind["criticas_altas_abiertas"]) == (1, 1)
    # ana (alerta + métrica), ext y beto; 2 + 3 eventos de alertas + 15 de métricas.
    assert (ind["usuarios_con_actividad"], ind["externos_con_actividad"], ind["eventos"]) == (3, 1, 20)
    assert ind["apps_detectadas"] == 3
    assert (ind["horas_medias_atencion"], ind["horas_medias_cierre"]) == (3.0, 10.0)
    assert (ind["equipos_con_actividad"], ind["equipos_fuera_de_tenable"]) == (2, 1)
    assert datos["cobertura_tenable"] == {"con_agente": 1, "por_escaneo": 0, "sin_inventario": 0, "fuera_de_tenable": 1}

    por_dia = {d["dia"]: d for d in datos["eventos_por_dia"]}
    assert len(por_dia) == 7  # todos los días del rango, aunque no tengan eventos
    assert por_dia["2026-10-18"] == {"dia": "2026-10-18", "baja": 0, "media": 10, "alta": 0, "critica": 2}
    assert por_dia["2026-10-14"]["alta"] == 1
    nuevas = {d["dia"]: d["cantidad"] for d in datos["alertas_nuevas_por_dia"]}
    assert (nuevas["2026-10-14"], nuevas["2026-10-18"]) == (1, 1)
    sensibles = {d["dia"]: d["actividades"] for d in datos["actividades_sensibles"]}
    assert sensibles["2026-10-18"] == {"Upload": 12, "Post": 0, "CallToolRequest": 0}
    assert sensibles["2026-10-14"]["CallToolRequest"] == 1

    assert datos["top_apps"][0] == {"nombre": "ChatGPT", "usuarios": 2, "eventos": 15}
    # El top de usuarios sólo mira reglas de alerta.
    assert [u["nombre"] for u in datos["top_usuarios"]] == ["ana@macro.com.ar", "ext@bmexternos.com.ar"]
    assert {p["poblacion"]: p["usuarios"] for p in datos["poblacion"]} == {"internos": 2, "externos": 1}
    assert datos["alertas_por_estado"] == {"nueva": 0, "en_revision": 1, "resuelta": 1, "descartada": 0}
    assert datos["software_tenable"] == [{"nombre": "Python", "cantidad": 1}]


def test_filtros(client: TestClient, auth_headers: dict[str, str], db: Session) -> None:
    reglas = _datos(db)

    def ind(**params):
        return client.get("/metricas/ia", params={**RANGO, **params}, headers=auth_headers).json()["indicadores"]

    assert ind(poblacion="externos")["usuarios_con_actividad"] == 1
    assert ind(poblacion="internos")["usuarios_con_actividad"] == 2
    assert ind(severidad="critica")["eventos"] == 3
    assert ind(regla_id=str(reglas["metrica"].id))["eventos"] == 15
    assert ind(desde="2026-09-01")["eventos"] == 20 + 1 + 99  # el rango amplio incluye lo viejo
    assert client.get("/metricas/ia", params={"poblacion": "otros"}, headers=auth_headers).status_code == 422


def test_requiere_autenticacion(client: TestClient) -> None:
    assert client.get("/metricas/ia").status_code == 401
    assert client.get("/metricas/ia/export.xlsx").status_code == 401


def test_export_excel(client: TestClient, auth_headers: dict[str, str], db: Session) -> None:
    _datos(db)

    r = client.get("/metricas/ia/export.xlsx", params={**RANGO, "severidad": "critica"}, headers=auth_headers)

    assert r.status_code == 200
    assert 'filename="seguridad-ia_2026-10-13_2026-10-19.xlsx"' in r.headers["content-disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Resumen", "Alertas", "Eventos", "Actividad diaria", "Equipos"]
    resumen = {fila[0]: fila[1] for fila in wb["Resumen"].iter_rows(values_only=True) if fila}
    assert (resumen["Severidad"], resumen["Eventos de IA"]) == ("critica", 3)
    alertas = list(wb["Alertas"].iter_rows(values_only=True))
    assert len(alertas) == 2 and alertas[1][4] == "ana@macro.com.ar"  # sólo la crítica dentro del rango
    eventos = list(wb["Eventos"].iter_rows(values_only=True))
    assert len(eventos) == 4
    # Lo que llega de Netskope se guarda como texto, nunca como fórmula.
    objeto = next(c for c in wb["Eventos"]["H"] if c.value and str(c.value).startswith("="))
    assert objeto.data_type == "s"


def test_export_csv(client: TestClient, auth_headers: dict[str, str], db: Session) -> None:
    _datos(db)

    r = client.get("/metricas/ia/export.csv", params={**RANGO, "tabla": "eventos"}, headers=auth_headers)

    assert r.status_code == 200 and r.content.startswith("﻿".encode())
    filas = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
    assert filas[0][:3] == ["Fecha", "Regla", "Severidad"]
    assert len(filas) == 1 + 5
    assert sum(f[7].startswith("'=") for f in filas[1:]) == 2  # fórmulas neutralizadas
    actividad = client.get("/metricas/ia/export.csv", params={**RANGO, "tabla": "actividad"}, headers=auth_headers)
    assert len(actividad.content.decode("utf-8-sig").strip().splitlines()) == 1 + 6
