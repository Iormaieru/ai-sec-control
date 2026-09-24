import io

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.catalog.models import Dominio, Pregunta
from app.documents.parsing import DOCX_CONTENT_TYPE


def _caso(client: TestClient, headers: dict, tipo: str = "ingresado", nombre: str = "Proyecto") -> dict:
    r = client.post(
        "/casos",
        json={"empresa_responsable": "ACME", "nombre_proyecto": nombre, "tipo": tipo},
        headers=headers,
    )
    assert r.status_code == 201, r.text
    return r.json()


def _backdate(db: Session, caso_id: str, when: str) -> None:
    db.execute(text("UPDATE casos SET created_at = :w WHERE id = :i"), {"w": when, "i": caso_id})
    db.commit()


def _metricas(client: TestClient, headers: dict, **params) -> dict:
    r = client.get("/metricas", params=params, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()


def test_metricas_sin_datos_devuelve_ceros(client: TestClient, auth_headers: dict[str, str]) -> None:
    m = _metricas(client, auth_headers)

    assert m["proyectos_ingresados"]["total"] == 0
    assert m["proyectos_ingresados"]["por_mes"] == []
    assert m["analisis"]["casos_analizados"] == 0
    assert m["modelados"]["cantidad"] == 0
    assert m["modelados"]["promedio_compliance_pct"] is None
    assert m["pentesting"]["total"] == 0
    assert m["pentesting"]["por_severidad"] == {"baja": 0, "media": 0, "alta": 0, "critica": 0}


def test_proyectos_ingresados_por_tipo_y_por_mes(
    client: TestClient, auth_headers: dict[str, str], db: Session
) -> None:
    viejo = _caso(client, auth_headers, "ingresado", "Viejo")
    _caso(client, auth_headers, "ingresado", "Nuevo 1")
    _caso(client, auth_headers, "candidato", "Nuevo 2")
    _backdate(db, viejo["id"], "2026-01-15 12:00:00+00")

    m = _metricas(client, auth_headers)["proyectos_ingresados"]

    assert m["total"] == 3
    assert m["por_tipo"] == {"candidato": 1, "ingresado": 2, "herramienta_tercero": 0}
    assert len(m["por_mes"]) == 2
    assert m["por_mes"][0] == {"periodo": "2026-01", "cantidad": 1}


def test_filtro_por_periodo(client: TestClient, auth_headers: dict[str, str], db: Session) -> None:
    viejo = _caso(client, auth_headers, nombre="Viejo")
    _caso(client, auth_headers, nombre="Nuevo")
    _backdate(db, viejo["id"], "2026-01-15 12:00:00+00")

    solo_enero = _metricas(client, auth_headers, desde="2026-01-01", hasta="2026-01-31")
    desde_junio = _metricas(client, auth_headers, desde="2026-06-01")

    assert solo_enero["proyectos_ingresados"]["total"] == 1
    assert desde_junio["proyectos_ingresados"]["total"] == 1


def test_filtro_por_tipo(client: TestClient, auth_headers: dict[str, str]) -> None:
    _caso(client, auth_headers, "ingresado")
    _caso(client, auth_headers, "herramienta_tercero")

    m = _metricas(client, auth_headers, tipo="herramienta_tercero")

    assert m["proyectos_ingresados"]["total"] == 1
    assert m["proyectos_ingresados"]["por_tipo"]["ingresado"] == 0


def test_analisis_cuenta_casos_con_documento_analizado(
    client: TestClient, auth_headers: dict[str, str], catalog_loaded: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    from app.llm.providers.mock_provider import MockProvider

    monkeypatch.setattr("app.documents.service.get_llm_provider", lambda: MockProvider())
    analizado = _caso(client, auth_headers, "candidato")
    _caso(client, auth_headers, "candidato")  # sin documento

    buf = io.BytesIO()
    doc = Document()
    doc.add_paragraph("Arquitectura")
    doc.save(buf)
    for _ in range(2):  # dos documentos sobre el mismo caso
        r = client.post(
            f"/casos/{analizado['id']}/documentos",
            files={"file": ("a.docx", buf.getvalue(), DOCX_CONTENT_TYPE)},
            headers=auth_headers,
        )
        assert r.status_code == 201, r.text

    a = _metricas(client, auth_headers)["analisis"]

    assert a["casos_analizados"] == 1
    assert a["documentos_analizados"] == 2
    assert a["por_tipo"]["candidato"] == 1


def test_modelados_incluye_solo_casos_con_respuestas_y_su_riesgo(
    client: TestClient, auth_headers: dict[str, str], db: Session, catalog_loaded: None
) -> None:
    caso = _caso(client, auth_headers, nombre="Con modelado")
    _caso(client, auth_headers, nombre="Sin modelado")

    # Mismo dataset real de Transparency & Accessibility usado en los golden tests: 75% en el dominio.
    client.post(f"/casos/{caso['id']}/preguntas", json={"dominio_codigo": "TAC"}, headers=auth_headers)
    dominio = db.query(Dominio).filter(Dominio.codigo == "TAC").one()
    preguntas = {p.numero: p for p in db.query(Pregunta).filter(Pregunta.dominio_id == dominio.id)}
    for numero, respuesta in {1: "SI", 2: "NO", 5: "SI"}.items():
        r = client.put(
            f"/casos/{caso['id']}/preguntas/{preguntas[numero].id}/respuesta",
            json={"respuesta": respuesta},
            headers=auth_headers,
        )
        assert r.status_code == 200

    m = _metricas(client, auth_headers)["modelados"]

    assert m["cantidad"] == 1
    assert m["casos"][0]["nombre_proyecto"] == "Con modelado"
    # global = 0.75 (TAC) * 0.06 (peso) — los otros 7 dominios sin evaluar aportan 0
    assert m["casos"][0]["compliance_pct"] == pytest.approx(0.045, abs=1e-9)
    assert m["casos"][0]["semaforo"] == "critico"
    assert m["por_semaforo"]["critico"] == 1
    assert m["promedio_compliance_pct"] == pytest.approx(0.045, abs=1e-9)


def test_pentesting_agrega_por_herramienta_y_severidad(
    client: TestClient, auth_headers: dict[str, str]
) -> None:
    caso_a = _caso(client, auth_headers, nombre="A")
    caso_b = _caso(client, auth_headers, nombre="B")
    burp = client.post("/pentesting/herramientas", json={"nombre": "Burp"}, headers=auth_headers).json()
    zap = client.post("/pentesting/herramientas", json={"nombre": "ZAP"}, headers=auth_headers).json()

    def resultado(caso, herramienta, severidad, fecha):
        r = client.post(
            f"/casos/{caso['id']}/pentests",
            json={"herramienta_id": herramienta["id"], "hallazgos": "x", "severidad": severidad, "fecha": fecha},
            headers=auth_headers,
        )
        assert r.status_code == 201

    resultado(caso_a, burp, "alta", "2026-03-01")
    resultado(caso_a, burp, "media", "2026-03-02")
    resultado(caso_b, zap, "alta", "2026-08-10")

    todos = _metricas(client, auth_headers)["pentesting"]
    assert todos["total"] == 3
    assert todos["casos_con_pentest"] == 2
    assert todos["por_herramienta"][0] == {"nombre": "Burp", "cantidad": 2}
    assert todos["por_severidad"] == {"baja": 0, "media": 1, "alta": 2, "critica": 0}

    solo_marzo = _metricas(client, auth_headers, desde="2026-03-01", hasta="2026-03-31")["pentesting"]
    assert solo_marzo["total"] == 2
    assert solo_marzo["casos_con_pentest"] == 1


def test_metricas_requiere_auth(client: TestClient) -> None:
    assert client.get("/metricas").status_code == 401


def _docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_export_docx_del_reporte(client: TestClient, auth_headers: dict[str, str]) -> None:
    caso = _caso(client, auth_headers, "ingresado", "Proyecto reportado")
    burp = client.post("/pentesting/herramientas", json={"nombre": "Burp"}, headers=auth_headers).json()
    client.post(
        f"/casos/{caso['id']}/pentests",
        json={"herramienta_id": burp["id"], "hallazgos": "x", "severidad": "alta", "fecha": "2026-03-01"},
        headers=auth_headers,
    )

    r = client.get(
        "/metricas/export/docx", params={"desde": "2026-01-01", "hasta": "2026-12-31"}, headers=auth_headers
    )

    assert r.status_code == 200
    assert "wordprocessingml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]
    texto = _docx_text(r.content)
    assert "Reporte de métricas" in texto
    assert "2026-01-01 a 2026-12-31" in texto
    assert "Burp" in texto
    assert "Ingresado" in texto


def test_export_docx_sin_datos_no_falla(client: TestClient, auth_headers: dict[str, str]) -> None:
    r = client.get("/metricas/export/docx", headers=auth_headers)

    assert r.status_code == 200
    assert "Sin pentestings en el período." in _docx_text(r.content)


def test_export_docx_requiere_auth(client: TestClient) -> None:
    assert client.get("/metricas/export/docx").status_code == 401
