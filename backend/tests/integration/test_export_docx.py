import io

import pytest
from docx import Document
from fastapi.testclient import TestClient


def _create_caso(client: TestClient, auth_headers: dict[str, str]) -> dict:
    response = client.post(
        "/casos",
        json={
            "gdld": "GDLD-42",
            "empresa_responsable": "ACME",
            "nombre_proyecto": "Chatbot de soporte",
            "tipo": "ingresado",
        },
        headers=auth_headers,
    )
    assert response.status_code == 201, response.text
    return response.json()


def _full_text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in document.paragraphs)


def test_export_docx_includes_caso_info_and_selected_questions(
    client: TestClient, auth_headers: dict[str, str], catalog_loaded: None
) -> None:
    caso = _create_caso(client, auth_headers)
    client.post(f"/casos/{caso['id']}/preguntas", json={"dominio_codigo": "TAC"}, headers=auth_headers)

    response = client.get(f"/casos/{caso['id']}/export/docx", headers=auth_headers)

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "attachment" in response.headers["content-disposition"]

    texto = _full_text(response.content)
    assert "Chatbot de soporte" in texto
    assert "ACME" in texto
    assert "GDLD 42" in texto
    # al menos una pregunta real de Transparency & Accessibility
    assert "explicarse de manera clara" in texto


def test_export_docx_includes_ai_instructions_when_present(
    client: TestClient, auth_headers: dict[str, str], monkeypatch: pytest.MonkeyPatch, catalog_loaded: None
) -> None:
    from app.llm.providers.mock_provider import MockProvider

    monkeypatch.setattr("app.documents.service.get_llm_provider", lambda: MockProvider())

    caso = _create_caso(client, auth_headers)
    docx_bytes = io.BytesIO()
    Document_ = __import__("docx").Document
    doc = Document_()
    doc.add_paragraph("Sistema que otorga créditos automáticamente.")
    doc.save(docx_bytes)

    client.post(
        f"/casos/{caso['id']}/documentos",
        files={
            "file": (
                "doc.docx",
                docx_bytes.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=auth_headers,
    )

    response = client.get(f"/casos/{caso['id']}/export/docx", headers=auth_headers)
    texto = _full_text(response.content)

    assert "Cómo responder:" in texto
    assert "How to answer:" in texto


def test_export_docx_requires_auth(client: TestClient, catalog_loaded: None) -> None:
    response = client.get("/casos/00000000-0000-0000-0000-000000000000/export/docx")
    assert response.status_code == 401
